import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule
from openpyxl.comments import Comment
from openpyxl.utils import get_column_letter as L

OUT = sys.argv[1]
TEST = len(sys.argv) > 2          # test mode: fill sample data to verify formulas

F = "Arial"
f_norm = Font(name=F, size=10)
f_bold = Font(name=F, size=10, bold=True)
f_hdr  = Font(name=F, size=10, bold=True, color="FFFFFF")
f_title= Font(name=F, size=14, bold=True, color="1F3864")
f_in   = Font(name=F, size=10, color="0000FF")
f_note = Font(name=F, size=9, italic=True, color="595959")
f_cmd  = Font(name="Courier New", size=9)
fill_h = PatternFill("solid", fgColor="1F3864")
fill_in= PatternFill("solid", fgColor="FFF9DB")
fill_ex= PatternFill("solid", fgColor="EDEDED")
fill_sec=PatternFill("solid", fgColor="D9E2F3")
thin = Side(style="thin", color="BFBFBF"); box = Border(left=thin,right=thin,top=thin,bottom=thin)
wrap = Alignment(wrap_text=True, vertical="top")
ctr  = Alignment(horizontal="center", vertical="center", wrap_text=True)

wb = Workbook()

def hdr(ws, row, labels, widths=None):
    for i, t in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=t)
        c.font, c.fill, c.alignment, c.border = f_hdr, fill_h, ctr, box
    ws.row_dimensions[row].height = 42
    if widths:
        for i, w in enumerate(widths, 1): ws.column_dimensions[L(i)].width = w

# ------------------------------------------------------------------ Start Here
ws = wb.active; ws.title = "Start Here"
ws.column_dimensions["A"].width = 26; ws.column_dimensions["B"].width = 95
ws["A1"] = "cSS semi-automatic grading pipeline - validation workbook"; ws["A1"].font = f_title
ws["A2"] = "Use this workbook to log every real case, record your review, and measure agreement between the pipeline and an expert reference."
ws["A2"].font = f_note
rows = [
 ("HOW TO USE", None),
 ("1. Prepare the scan", "Need the SWI series (processed SWI or magnitude - NOT the minIP and NOT the phase map), plus the 3-D T1 and FLAIR if available (see PIPELINE.md). DICOM: convert with dcm2niix (see Commands). Use a code like P006 - never a name or MRN."),
 ("2. Run the pipeline", "Terminal:  run_css.sh P006 /path/to/swi.nii --t1 /path/to/t1.nii --flair /path/to/flair.nii [--recon]   (anatomy, detects, ranks candidates, opens freeview)."),
 ("3. Export candidates", "Terminal:  python ~/css_project/scripts/export_review.py P006 40 | pbcopy   then paste into 'Candidate Review' at column A of the next empty row."),
 ("4. Review in freeview", "Jump to each candidate (vox_i vox_j vox_k). In 'Candidate Review' choose a Reader call: cSS / Vein / Artifact / Near ICH / Unsure. Use 'Near ICH' for siderosis connected to a lobar haemorrhage - it is reported but NOT scored (standard convention). IDs build automatically."),
 ("5. Score", "Copy the command from 'Case Log' column V into Terminal. Type the printed score (0-4) and foci into Case Log columns I and J. Mark column W if the scan shows a lobar ICH."),
 ("6. Compare to the expert", "Enter the expert reference in Case Log columns G (cSS present 1/0) and H (multifocality 0-4). 'Validation' updates automatically."),
 ("COLOUR LEGEND", None),
 ("Blue text, pale-yellow fill", "Cells YOU type into."),
 ("Black text", "Formulas - do not overwrite."),
 ("Grey row 6", "Example row showing the expected format. It is NOT counted in the statistics (statistics start at row 7)."),
 ("SHEETS", None),
 ("Case Log", "One row per patient: scan details, expert reference, pipeline result, agreement flags (up to 200 patients)."),
 ("Candidate Review", "One row per candidate you reviewed (up to 2,000 rows). Builds the accepted-ID list."),
 ("Validation", "Sensitivity, specificity, kappa (presence), weighted kappa (0-4 score), exact/within-1 agreement, review depth. All formulas."),
 ("Commands", "Copy-paste terminal commands for every step, including DICOM conversion."),
 ("Parameters", "Final detector settings, reasons, and the synthetic-validation results they came from."),
 ("READ BEFORE REAL TESTING", None),
 ("Data governance", "Keep scans on Mayo-approved storage. Use study codes only in this workbook. Confirm with Dr. Lin / Mayo IT that the data may sit on this Mac."),
 ("Known limits", "Tuned on one skull-stripped SWI slab (2 mm slices) plus synthetic lesions; real-case accuracy is unknown until you fill this workbook. Faint cSS (<~70% of cortex intensity) is often missed. Medial cSS near the midline is down-ranked. Skull-base slices are artifact-prone. GRE (4 mm) is untested."),
 ("Blinding", "For a clean validation, review candidates BEFORE looking at the expert rating, and have the expert rate independently."),
]
r = 4
for a, b in rows:
    if b is None:
        c = ws.cell(row=r, column=1, value=a); c.font = f_bold; c.fill = fill_sec
        ws.cell(row=r, column=2).fill = fill_sec
    else:
        ws.cell(row=r, column=1, value=a).font = f_bold
        ws.cell(row=r, column=2, value=b).font = f_norm
        ws.cell(row=r, column=1).alignment = wrap; ws.cell(row=r, column=2).alignment = wrap
        ws.row_dimensions[r].height = 15 * (1 + len(b) // 100)
    r += 1
ws["A12"].font = f_in; ws["A12"].fill = fill_in
ws["A13"].font = f_norm
ws["A14"].fill = fill_ex

# ------------------------------------------------------------------ Candidate Review
cr = wb.create_sheet("Candidate Review")
CR0, CR1 = 7, 2006
crh = ["Subject ID","cand_id","hemi","region","label","volume_mm3","n_slices","elongation",
       "surface_contact","darkness_z","artifact_zone","midline_zone","score","vox_i","vox_j","vox_k",
       "Reader call (cSS / Vein / Artifact / Near ICH / Unsure)","Accept (1/0)","Accepted IDs so far","Near-ICH IDs so far","Notes"]
cr["A1"] = "Candidate Review"; cr["A1"].font = f_title
cr["A2"] = "Paste the export (columns A-P) from export_review.py. Then pick a Reader call in column Q for each candidate you looked at. Keep each subject's rows together."
cr["A2"].font = f_note
cr["A3"] = "Blue = type/paste. Black = formula. Row 6 is an example and is not counted."; cr["A3"].font = f_note
hdr(cr, 5, crh, [12,8,6,30,8,11,9,10,12,11,10,10,8,7,7,7,24,9,16,16,32])
ex = ["EXAMPLE-01",2,"L","ctx-lh-superiortemporal",1030,73.0,9,5.1,0.70,4.5,0,0,18.41,361,216,11,"cSS",None,None,None,"Gyriform band, both banks of sulcus"]
for r in [6] + list(range(CR0, CR1 + 1)):
    for c in range(1, 22):
        cell = cr.cell(row=r, column=c)
        cell.font = f_norm; cell.border = box
        if c in (1,2,3,4,5,6,7,8,9,10,11,12,13,14,15,16,17,21):
            cell.font = f_in; cell.fill = fill_in
    cr.cell(row=r, column=18, value=f'=IF(Q{r}="","",IF(Q{r}="cSS",1,0))')
    pv = f'IF(A{r}=A{r-1},S{r-1},"")'
    cr.cell(row=r, column=19, value=f'=IF(A{r}="","",IF(R{r}=1,IF({pv}="","",{pv}&",")&B{r},{pv}))')
    pi = f'IF(A{r}=A{r-1},T{r-1},"")'
    cr.cell(row=r, column=20, value=f'=IF(A{r}="","",IF(Q{r}="Near ICH",IF({pi}="","",{pi}&",")&B{r},{pi}))')
for c, v in enumerate(ex, 1):
    if v is not None: cr.cell(row=6, column=c, value=v)
for c in range(1, 22):
    cr.cell(row=6, column=c).fill = fill_ex
dv = DataValidation(type="list", formula1='"cSS,Vein,Artifact,Near ICH,Unsure"', allow_blank=True)
cr.add_data_validation(dv); dv.add(f"Q6:Q{CR1}")
cr.freeze_panes = "E6"

# ------------------------------------------------------------------ Case Log
cl = wb.create_sheet("Case Log", 1)
CL0, CL1 = 7, 206
clh = ["Subject ID","Cohort / source","Sequence","Field (T)","Slice thickness (mm)","Run date",
       "EXPERT: cSS present (1/0)","EXPERT: multifocality (0-4)","PIPELINE: multifocality (0-4)","PIPELINE: n foci",
       "Pipeline: cSS present","Presence agrees (1/0)","Score diff (pipeline - expert)","Abs diff","Exact score match","Within +/-1",
       "Rank of first true cSS candidate","Review time (min)","Accepted candidate IDs","Accepted count","Near-ICH candidate IDs (not scored)","Command to score (run after review)","Lobar ICH on scan (1/0)","Notes"]
cl["A1"] = "Case Log"; cl["A1"].font = f_title
cl["A2"] = "One row per patient. Blue/yellow = you type. Black = formula. Row 6 is an example and is not counted in the statistics."; cl["A2"].font = f_note
cl["A3"] = "Columns I-J: copy from the mark_css.py printout (score and total foci). Columns G-H: expert reference (enter independently)."; cl["A3"].font = f_note
hdr(cl, 5, clh, [13,18,10,9,12,12,12,13,13,10,11,11,13,8,9,9,14,10,18,9,16,64,11,32])
inputs = {1,2,3,4,5,6,7,8,9,10,17,18,23,24}
CRrng = lambda col: f"'Candidate Review'!${col}${CR0}:${col}${CR1}"
for r in [6] + list(range(CL0, CL1 + 1)):
    for c in range(1, 25):
        cell = cl.cell(row=r, column=c); cell.border = box; cell.font = f_norm
        if c in inputs: cell.font = f_in; cell.fill = fill_in
    cl.cell(row=r, column=11, value=f'=IF(I{r}="","",IF(I{r}>0,1,0))')
    cl.cell(row=r, column=12, value=f'=IF(OR(G{r}="",K{r}=""),"",IF(G{r}=K{r},1,0))')
    cl.cell(row=r, column=13, value=f'=IF(OR(H{r}="",I{r}=""),"",I{r}-H{r})')
    cl.cell(row=r, column=14, value=f'=IF(M{r}="","",ABS(M{r}))')
    cl.cell(row=r, column=15, value=f'=IF(M{r}="","",IF(M{r}=0,1,0))')
    cl.cell(row=r, column=16, value=f'=IF(M{r}="","",IF(N{r}<=1,1,0))')
    if r != 6:
        cl.cell(row=r, column=19, value=(
            f'=IF(A{r}="","",IF(COUNTIF({CRrng("A")},A{r})=0,"",'
            f'INDEX({CRrng("S")},SUMPRODUCT(MAX(({CRrng("A")}=A{r})*(ROW({CRrng("A")})-{CR0-1}))))))'))
    cl.cell(row=r, column=20, value=f'=IF(A{r}="","",IF(S{r}="",0,LEN(S{r})-LEN(SUBSTITUTE(S{r},",",""))+1))')
    if r != 6:
        cl.cell(row=r, column=21, value=(
            f'=IF(A{r}="","",IF(COUNTIF({CRrng("A")},A{r})=0,"",'
            f'INDEX({CRrng("T")},SUMPRODUCT(MAX(({CRrng("A")}=A{r})*(ROW({CRrng("A")})-{CR0-1}))))))'))
    cl.cell(row=r, column=22, value=f'=IF(A{r}="","","python ~/css_project/scripts/mark_css.py "&A{r}&" "&IF(S{r}="","none",S{r})&IF(U{r}="",""," --ich "&U{r}))')
    cl.cell(row=r, column=22).font = f_cmd
    cl.cell(row=r, column=19).font = f_norm
exv = ["EXAMPLE-01","Mayo CAA cohort","SWI",3,2,"2026-10-02",1,3,2,3,None,None,None,None,None,None,4,6,"2,5",None,"7",None,1,"Example only - not counted"]
for c, v in enumerate(exv, 1):
    if v is not None: cl.cell(row=6, column=c, value=v)
for c in range(1, 25): cl.cell(row=6, column=c).fill = fill_ex
cl.cell(row=6, column=19).font = f_in; cl.cell(row=6, column=21).font = f_in
for rng, dvv in [(f"C6:C{CL1}", '"SWI,minIP,T2* GRE,Other"'), (f"G6:G{CL1}", '"0,1"'), (f"W6:W{CL1}", '"0,1"')]:
    d = DataValidation(type="list", formula1=dvv, allow_blank=True); cl.add_data_validation(d); d.add(rng)
d = DataValidation(type="whole", operator="between", formula1="0", formula2="4", allow_blank=True)
d.error = "Score must be a whole number 0-4"; cl.add_data_validation(d); d.add(f"H6:I{CL1}")
cl.conditional_formatting.add(f"L{CL0}:L{CL1}", CellIsRule(operator="equal", formula=["0"], fill=PatternFill("solid", bgColor="F8CBAD", fgColor="F8CBAD")))
cl.conditional_formatting.add(f"L{CL0}:L{CL1}", CellIsRule(operator="equal", formula=["1"], fill=PatternFill("solid", bgColor="C6EFCE", fgColor="C6EFCE")))
cl.freeze_panes = "B6"
cl["I5"].comment = Comment("Copy the 'cSS multifocality score: x/4' printed by mark_css.py", "Pipeline")
cl["H5"].comment = Comment("Expert/reference total multifocality score 0-4 (sum of the two hemisphere scores).", "Pipeline")
cl["Q5"].comment = Comment("Position in the ranked candidate list of the first candidate you accepted as true cSS. Used for review-depth statistics.", "Pipeline")

# ------------------------------------------------------------------ Validation
va = wb.create_sheet("Validation", 2)
va.column_dimensions["A"].width = 50
for c in "BCDEFG": va.column_dimensions[c].width = 13
va.column_dimensions["H"].width = 60
va["A1"] = "Validation - pipeline versus expert reference"; va["A1"].font = f_title
va["A2"] = "All cells are formulas reading 'Case Log' rows 7-206 (the example row is excluded)."; va["A2"].font = f_note
G = f"'Case Log'!$G${CL0}:$G${CL1}"; H = f"'Case Log'!$H${CL0}:$H${CL1}"; I = f"'Case Log'!$I${CL0}:$I${CL1}"
K = f"'Case Log'!$K${CL0}:$K${CL1}"; O = f"'Case Log'!$O${CL0}:$O${CL1}"; P = f"'Case Log'!$P${CL0}:$P${CL1}"
N = f"'Case Log'!$N${CL0}:$N${CL1}"; Q = f"'Case Log'!$Q${CL0}:$Q${CL1}"; R = f"'Case Log'!$R${CL0}:$R${CL1}"

def sec(r, t):
    for c in range(1, 9): va.cell(row=r, column=c).fill = fill_sec
    va.cell(row=r, column=1, value=t).font = f_bold
def line(r, label, formula, fmt=None, note=None):
    va.cell(row=r, column=1, value=label).font = f_norm
    c = va.cell(row=r, column=2, value=formula); c.font = f_norm; c.border = box
    if fmt: c.number_format = fmt
    if note: va.cell(row=r, column=8, value=note).font = f_note

sec(4, "A. Presence of cSS (case level)")
line(5, "Cases with expert and pipeline result", f'=COUNTIFS({G},">=0",{K},">=0")', "0")
line(6, "Expert positive", f'=COUNTIFS({G},1,{K},">=0")', "0")
line(7, "Expert negative", f'=COUNTIFS({G},0,{K},">=0")', "0")
line(8, "True positive (both say cSS)", f'=COUNTIFS({G},1,{K},1)', "0")
line(9, "False negative (pipeline missed)", f'=COUNTIFS({G},1,{K},0)', "0")
line(10, "False positive (pipeline score > 0, expert none)", f'=COUNTIFS({G},0,{K},1)', "0")
line(11, "True negative", f'=COUNTIFS({G},0,{K},0)', "0")
line(12, "Sensitivity", '=IF(B8+B9=0,"n/a",B8/(B8+B9))', "0%")
line(13, "Specificity", '=IF(B11+B10=0,"n/a",B11/(B11+B10))', "0%")
line(14, "Positive predictive value", '=IF(B8+B10=0,"n/a",B8/(B8+B10))', "0%")
line(15, "Negative predictive value", '=IF(B11+B9=0,"n/a",B11/(B11+B9))', "0%")
line(16, "Accuracy", '=IF(B5=0,"n/a",(B8+B11)/B5)', "0%")
line(17, "Chance-expected agreement", '=IF(B5=0,"n/a",((B8+B9)*(B8+B10)+(B10+B11)*(B9+B11))/B5^2)', "0.00")
line(18, "Cohen's kappa (presence)", '=IF(B5=0,"n/a",IF(B17=1,"n/a",(B16-B17)/(1-B17)))', "0.00",
     "Conventional bands (Landis & Koch 1977): <0.20 slight, 0.21-0.40 fair, 0.41-0.60 moderate, 0.61-0.80 substantial, >0.80 almost perfect.")
sec(20, "B. Multifocality score 0-4 (agreement)")
line(21, "Cases with both scores", f'=COUNTIFS({H},">=0",{I},">=0")', "0")
line(22, "Exact match", f'=IF(B21=0,"n/a",SUM({O})/B21)', "0%")
line(23, "Within +/-1 point", f'=IF(B21=0,"n/a",SUM({P})/B21)', "0%")
line(24, "Mean absolute difference (points)", f'=IF(B21=0,"n/a",SUM({N})/B21)', "0.00")
line(25, "Quadratic-weighted kappa", '=IF(B21=0,"n/a",IF(B50=1,"n/a",(B49-B50)/(1-B50)))', "0.00",
     "Computed from the 5x5 table below (weights = 1 - (i-j)^2/16). Needs several cases per score level to be stable.")
sec(27, "C. Review depth (expert-positive cases with a rank entered)")
line(28, "Cases with rank entered", f'=COUNTIFS({G},1,{Q},">0")', "0")
line(29, "First true cSS candidate within top 10", f'=IF(B28=0,"n/a",COUNTIFS({G},1,{Q},">0",{Q},"<=10")/B28)', "0%")
line(30, "...within top 20", f'=IF(B28=0,"n/a",COUNTIFS({G},1,{Q},">0",{Q},"<=20")/B28)', "0%")
line(31, "...within top 45", f'=IF(B28=0,"n/a",COUNTIFS({G},1,{Q},">0",{Q},"<=45")/B28)', "0%",
     "Synthetic testing suggested reviewing the top ~45 candidates catches strong and medium cSS. Compare with your real result.")
line(32, "Mean rank of first true cSS candidate", f'=IF(B28=0,"n/a",AVERAGEIFS({Q},{G},1,{Q},">0"))', "0.0")
line(33, "Mean review time per case (min)", f'=IFERROR(AVERAGE({R}),"n/a")', "0.0")

sec(35, "D. 5x5 table for weighted kappa (rows = expert score, columns = pipeline score)")
va["A36"] = "Expert \\ Pipeline"; va["A36"].font = f_bold
for j in range(5):
    c = va.cell(row=36, column=2 + j, value=j); c.font = f_bold; c.alignment = ctr
va.cell(row=36, column=7, value="Total").font = f_bold
for i in range(5):
    r = 37 + i
    va.cell(row=r, column=1, value=i).font = f_bold; va.cell(row=r, column=1).alignment = Alignment(horizontal="right")
    for j in range(5):
        c = va.cell(row=r, column=2 + j, value=f'=COUNTIFS({H},$A{r},{I},{L(2+j)}$36)'); c.font = f_norm; c.border = box
    va.cell(row=r, column=7, value=f'=SUM(B{r}:F{r})').font = f_bold
va["A42"] = "Total"; va["A42"].font = f_bold
for j in range(5):
    va.cell(row=42, column=2 + j, value=f'=SUM({L(2+j)}37:{L(2+j)}41)').font = f_bold
va["G42"] = "=SUM(G37:G41)"; va["G42"].font = f_bold
va["A44"] = "Weights (quadratic)"; va["A44"].font = f_note
for i in range(5):
    r = 45 + i if False else None
# weights and expected blocks placed to the right (columns J-O / Q-V) to keep B49/B50 fixed
va["J36"] = "Weight"; va["J36"].font = f_bold
va["Q36"] = "Expected"; va["Q36"].font = f_bold
for j in range(5):
    va.cell(row=36, column=11 + j, value=j).font = f_bold
    va.cell(row=36, column=18 + j, value=j).font = f_bold
for i in range(5):
    r = 37 + i
    va.cell(row=r, column=10, value=i).font = f_bold
    va.cell(row=r, column=17, value=i).font = f_bold
    for j in range(5):
        va.cell(row=r, column=11 + j, value=f'=1-(($J{r}-{L(11+j)}$36)^2)/16').font = f_norm
        va.cell(row=r, column=18 + j, value=f'=IF($G$42=0,0,$G{r}*{L(2+j)}$42/$G$42)').font = f_norm
        va.cell(row=r, column=11 + j).number_format = "0.00"; va.cell(row=r, column=18 + j).number_format = "0.00"
va["A49"] = "Observed weighted agreement"; va["A49"].font = f_norm
va["B49"] = '=IF(G42=0,0,SUMPRODUCT(K37:O41,B37:F41)/G42)'; va["B49"].number_format = "0.000"; va["B49"].border = box
va["A50"] = "Chance-expected weighted agreement"; va["A50"].font = f_norm
va["B50"] = '=IF(G42=0,0,SUMPRODUCT(K37:O41,R37:V41)/G42)'; va["B50"].number_format = "0.000"; va["B50"].border = box
for c in range(10, 23): va.column_dimensions[L(c)].width = 7
va.column_dimensions["I"].width = 3

# ------------------------------------------------------------------ Commands
cm = wb.create_sheet("Commands")
cm["A1"] = "Terminal commands (Mac, FreeSurfer 8.2.0, conda env 'css')"; cm["A1"].font = f_title
hdr(cm, 3, ["Step", "What it does", "Command (replace P006 and the paths)"], [8, 46, 120])
cmds = [
 ("0", "Open Terminal. Prompt must start with (css).", "conda activate css"),
 ("1a", "DICOM -> NIfTI (any MRI). Lists all series.", "mkdir -p ~/css_project/nifti/P006 && dcm2niix -z n -f \"%s_%d\" -o ~/css_project/nifti/P006 /path/to/dicom_folder && ls ~/css_project/nifti/P006"),
 ("1b", "Choose the SWI (processed SWI or magnitude) .nii - NOT the minIP, NOT phase (_ph). Also note the 3-D T1 and FLAIR files.", "ls ~/css_project/nifti/P006"),
 ("1c", "Check size/orientation (needs full axial brain slab, ideally <=2 mm slices).", "mri_info ~/css_project/nifti/P006/YOURFILE.nii | head -12"),
 ("2", "RUN: import, T1/FLAIR anatomy, align, detect, rank, open freeview (add --recon for sulcal scoring).", "run_css.sh P006 ~/css_project/nifti/P006/SWI.nii --t1 ~/css_project/nifti/P006/T1.nii --flair ~/css_project/nifti/P006/FLAIR.nii"),
 ("3", "Export top-40 candidates to the clipboard, then paste into 'Candidate Review' (column A, next empty row).", "python ~/css_project/scripts/export_review.py P006 40 | pbcopy"),
 ("4", "Open the scan again later (candidates in colour).", "freeview -v ~/css_project/data/P006_swi.nii ~/css_project/review/P006_candidates.nii.gz:colormap=lut:opacity=0.6 &"),
 ("5", "Score after review (copy the ready-made line from Case Log column V). --ich lists siderosis connected to a lobar ICH (reported, not scored).", "python ~/css_project/scripts/mark_css.py P006 2,5,9 --ich 7      # or:  ... P006 none"),
 ("6", "Open the all-patients score table.", "open ~/css_project/review/css_scores.csv"),
 ("7", "Back up scripts + results (change the destination).", "cp -R ~/css_project/scripts ~/css_project/review ~/Documents/css_backup/"),
 ("opt", "Higher sensitivity mode for suspected FAINT cSS (more candidates).", "python ~/css_project/scripts/detect_css.py P006 -2.0 85 3"),
 ("opt", "Batch: run several patients (edit the list).", "run_all.sh P006 P007 P008    (each needs data/<ID>_swi.nii; logs in ~/css_project/logs)"),
]
for i, (a, b, c) in enumerate(cmds, 4):
    cm.cell(row=i, column=1, value=a).font = f_bold
    cm.cell(row=i, column=2, value=b).font = f_norm
    cm.cell(row=i, column=3, value=c).font = f_cmd
    for k in (1, 2, 3):
        cm.cell(row=i, column=k).alignment = wrap; cm.cell(row=i, column=k).border = box
    cm.row_dimensions[i].height = 30
r0 = 4 + len(cmds) + 1
cm.cell(row=r0, column=1, value="One-time helper script (paste once into Terminal) - exports candidates in a form that pastes straight into Excel:").font = f_bold
helper = ('cat > ~/css_project/scripts/export_review.py << \'EOF\'\n'
 'import sys, pandas as pd\n'
 'subj = sys.argv[1]; n = int(sys.argv[2]) if len(sys.argv) > 2 else 40\n'
 'df = pd.read_csv(f"/Users/tarabeah.selim/css_project/review/{subj}_candidates.csv").head(n)\n'
 'cols = ["cand_id","hemi","region","label","volume_mm3","n_slices","elongation","surface_contact","darkness_z","artifact_zone","midline_zone","score","vox_i","vox_j","vox_k"]\n'
 'for _, r in df.iterrows():\n'
 '    print(subj + "\\t" + "\\t".join(str(r[c]) for c in cols))\n'
 'EOF')
c = cm.cell(row=r0 + 1, column=3, value=helper); c.font = f_cmd; c.alignment = wrap
cm.merge_cells(start_row=r0 + 1, start_column=2, end_row=r0 + 1, end_column=3)
cm.cell(row=r0 + 1, column=2, value=helper).font = f_cmd
cm.cell(row=r0 + 1, column=2).alignment = wrap
cm.row_dimensions[r0 + 1].height = 150

# ------------------------------------------------------------------ Parameters
pm = wb.create_sheet("Parameters")
pm["A1"] = "Final detector settings and where they came from"; pm["A1"].font = f_title
hdr(pm, 3, ["Setting", "Value", "Why", "Evidence"], [34, 14, 70, 70])
prm = [
 ("Segmentation", "SynthSeg-robust 2.0 (--parc)", "Contrast-agnostic cortex/CSF labels directly on SWI; no T1 needed.", "Run on P005; labels followed the cortical ribbon."),
 ("Coverage mask", "SWI > 0", "SynthSeg imagines anatomy outside the slab; labels are removed where SWI has no signal.", "Removed ~67,000 voxels (2.3%) on P005."),
 ("Edge rim excluded", "3 mm (3D, mm units)", "Skull-stripped edge looks like a dark ridge; 2D rim was too thin.", "P005: 3 mm gave 8/8 synthetic lesions, worst rank #15; 2 mm gave 7/8, worst rank #28."),
 ("Ridge filter percentile", "85", "Ridge strength scales with contrast; 95th percentile was set by veins and rejected cSS-like bands.", "At 95: 1/8 lesions; at 85: 8/8 lesions with ~73% of voxels intact."),
 ("Darkness cutoff (z)", "-2.5", "Robust z vs cortex median/MAD. -2.0 is an optional high-sensitivity mode.", "Synthetic (P005, 3 layouts x 3 depths): -2.5 found 22/24 strong, 20/24 medium, 11/24 faint; -2.0 found 23/24, 20/24, 15/24 but ranks reached #85-86 and clean-scan candidates rose 83 -> 106."),
 ("Min size / slices / elongation", "10 mm3 / 2 slices / 2.5", "Removes single-slice specks and round blobs (microbleeds).", "Cut raw candidate list from 351 to ~100 on P005."),
 ("Midline handling", "Split by hemisphere; <5 mm from midline = half score", "Interhemispheric fissure merged lesions across hemispheres and topped the list.", "Fixed hemisphere assignment (left 5 / right 3 on synthetic test); synthetic medial lesion fell to rank #38, other 7 in top 11."),
 ("Artifact zones", "orbitofrontal, entorhinal, temporal pole, parahippocampal, fusiform, inferior temporal: half score", "Air-bone susceptibility at skull base.", "Flagged, not deleted."),
 ("Same-sulcus merge distance", "3 mm", "Pieces of one band count as one focus.", "Design choice - not validated on real cases."),
 ("Adjacent-sulcus distance", "10 mm", "Foci within 10 mm = one contiguous cluster.", "Design choice - not validated on real cases."),
 ("Scoring rule (per hemisphere)", "0 none; 1 = one cluster of <=3 foci; 2 = >1 cluster or >3 foci. Total 0-4", "Approximates the published multifocality rule using physical distance because SWI-only segmentation gives gyri, not sulci.", "Verified on synthetic lesions: 4/4 disseminated."),
 ("Focal / disseminated", "focal <=3 foci; disseminated >3", "Conventional categories.", "Definition used in the literature."),
 ("Near-ICH siderosis", "Excluded from the 0-4 score; count and volume reported", "cSS potentially connected to a lobar ICH is not included in multifocality ratings.", "Convention as applied by van Harten et al. 2023 (NeuroImage: Clinical 38:103447)."),
 ("Vein penalty (v3)", "cortex_frac < 0.35 -> 0.6x score", "Veins lie mid-sulcus; cSS lines the cortex.", "Phantom test only (cSS 0.92 vs veins 0.35). Validate on real cases."),
 ("Surface depth profile (v3)", "surface_gradient recorded, not used for ranking", "cSS iron is densest in the outermost cortex (MGH histology).", "Phantom only; partly built into the synthetic test, so real cases decide."),
]
for i, row in enumerate(prm, 4):
    for k, v in enumerate(row, 1):
        c = pm.cell(row=i, column=k, value=v); c.font = f_bold if k == 1 else f_norm; c.alignment = wrap; c.border = box
    pm.row_dimensions[i].height = 62
pm.cell(row=4 + len(prm) + 1, column=1, value="Source of all numbers above: this project's own synthetic-lesion tests on case P005 (October 2026), run before the final midline changes. Re-run ~/css_project/scripts/stress_test.sh to refresh them.").font = f_note

# ------------------------------------------------------------------ test data
if TEST:
    data = [  # id, G, H, I, J, Q
        ("T01",1,3,3,4,2),("T02",1,2,1,2,5),("T03",1,4,4,6,1),("T04",0,0,0,0,None),
        ("T05",0,0,1,1,None),("T06",1,1,0,0,None),("T07",0,0,0,0,None),("T08",1,2,2,3,12)]
    for k, (a, g, h, i_, j, q) in enumerate(data):
        r = CL0 + k
        cl.cell(row=r, column=1, value=a); cl.cell(row=r, column=7, value=g); cl.cell(row=r, column=8, value=h)
        cl.cell(row=r, column=9, value=i_); cl.cell(row=r, column=10, value=j)
        if q: cl.cell(row=r, column=17, value=q)
        cl.cell(row=r, column=18, value=5 + k)
    cand = [("T01",1,"cSS"),("T01",2,"Vein"),("T01",3,"cSS"),("T02",1,"Vein"),("T02",4,"cSS"),("T02",5,"cSS"),("T03",1,"Artifact")]
    for k, (a, cid, call) in enumerate(cand):
        r = CR0 + k
        cr.cell(row=r, column=1, value=a); cr.cell(row=r, column=2, value=cid); cr.cell(row=r, column=17, value=call)

wb.save(OUT)
print("saved", OUT)
